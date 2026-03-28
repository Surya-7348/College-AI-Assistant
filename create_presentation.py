from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define color scheme
DARK_BLUE = RGBColor(15, 32, 39)
LIGHT_BLUE = RGBColor(42, 82, 152)
ACCENT_BLUE = RGBColor(44, 83, 100)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(245, 247, 250)
TEXT_DARK = RGBColor(51, 51, 51)

def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Header bar
    header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = DARK_BLUE
    header_shape.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.8))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, point in enumerate(content_points):
        if i > 0:
            text_frame.add_paragraph()
        p = text_frame.paragraphs[i]
        p.text = point
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(12)
        p.space_after = Pt(12)
        p.level = 0
    
    return slide

# Slide 1: Title Slide
add_title_slide(prs, "🎓 College AI Assistant", "Intelligent Learning Platform for Students & Educators\nBuilt with Google Gemini AI")

# Slide 2: Problem Statement
add_content_slide(prs, "The Challenge", [
    "📚 Students need instant answers to college-related questions",
    "⏰ Traditional resources are time-consuming and not always available 24/7",
    "🤝 Educators lack modern tools for student engagement",
    "💡 Growing demand for AI-powered educational solutions"
])

# Slide 3: Solution Overview
add_content_slide(prs, "Our Solution", [
    "✨ An intelligent AI chatbot powered by Google Gemini 2.5 Flash",
    "🎯 Real-time Q&A for college-related topics",
    "📱 Clean, professional web interface",
    "⚡ Fast, concise, and professional responses",
    "🔄 Conversation history & context awareness"
])

# Slide 4: Project Objectives
add_content_slide(prs, "Project Objectives", [
    "🎯 Provide 24/7 AI assistance for college students",
    "📖 Answer questions about courses, admissions, careers & more",
    "🏆 Deliver professional, educational responses",
    "🎨 Create an intuitive and modern user interface",
    "🚀 Deploy as a full-stack web application"
])

# Slide 5: Technology Stack
add_content_slide(prs, "Technology Stack", [
    "🖥️ Frontend: HTML, CSS, JavaScript (Vanilla JS)",
    "🔧 Backend: Node.js, Express.js",
    "🤖 AI Engine: Google Gemini 2.5 Flash API",
    "📦 Package Manager: npm",
    "🎯 Deployment Ready: Heroku, Vercel, Railway"
])

# Slide 6: Architecture Overview
add_content_slide(prs, "System Architecture", [
    "📍 Client-Side: Browser with chat interface",
    "🔗 Communication: REST API (HTTP POST requests)",
    "⚙️ Server: Node.js Express server listening on port 3003",
    "🤖 AI Processing: Google Generative AI SDK",
    "💾 Data Flow: Messages → Server → Gemini API → Response"
])

# Slide 7: Frontend Features
add_content_slide(prs, "Frontend Highlights", [
    "✨ Modern, professional UI with dark blue gradient theme",
    "💬 Real-time chat interface with message bubbles",
    "✓ Live server connection status indicator",
    "⌨️ Keyboard support (Enter to send, Shift+Enter for new line)",
    "📱 Responsive design for all screen sizes",
    "🎨 Smooth animations and professional typography"
])

# Slide 8: Backend Features
add_content_slide(prs, "Backend Capabilities", [
    "🔐 CORS enabled for cross-origin requests",
    "📝 Request validation and error handling",
    "🤖 Gemini API integration with system prompts",
    "💬 Conversation history management",
    "📊 Professional response formatting",
    "🐛 Detailed error logging for debugging"
])

# Slide 9: AI Integration
add_content_slide(prs, "Gemini AI Integration", [
    "🧠 Model: Google Gemini 2.5 Flash (latest)",
    "📋 System Prompt: Ensures professional, concise responses",
    "---",
    "Key Constraints:",
    "  • Limit responses to 2-3 paragraphs",
    "  • Professional & educational tone",
    "  • Direct and actionable information"
])

# Slide 10: User Interface Demo
add_content_slide(prs, "User Interface", [
    "🎨 Header: College AI Assistant branding",
    "💬 Chat Area: Clean conversation display",
    "📤 Input Area: Message box with Send button",
    "✓ Status: Real-time server connection indicator",
    "🎯 UX: Intuitive and minimalist design",
    "📱 Mobile-Friendly: Works on all devices"
])

# Slide 11: Chat Capabilities
add_content_slide(prs, "What Can It Do?", [
    "❓ Answer college admission questions",
    "📚 Explain academic subjects & topics",
    "🎓 Provide career guidance & advice",
    "🏫 Share information about institutions",
    "📝 Help with project & assignment ideas",
    "💼 Discuss professional development"
])

# Slide 12: Response Quality
add_content_slide(prs, "Response Quality Features", [
    "✅ Professionally formatted answers",
    "📖 Concise but comprehensive information",
    "🎯 Direct answers to specific questions",
    "📚 Educational & informative content",
    "⚡ Instant responses (no loading delays)",
    "🔄 Context-aware with conversation history"
])

# Slide 13: Use Cases for Students
add_content_slide(prs, "Student Use Cases", [
    "🎓 Quick answers between classes",
    "📚 Study assistance & explanations",
    "🏫 College selection & admission help",
    "💼 Career planning & internship advice",
    "🤔 Help with assignments & projects",
    "🕐 24/7 availability for last-minute questions"
])

# Slide 14: Use Cases for Educators
add_content_slide(prs, "Educator Use Cases", [
    "📖 Supplement classroom teaching",
    "🎯 Student engagement tool",
    "⏰ Reduce repetitive Q&A in classes",
    "📝 Assign chatbot-assisted learning tasks",
    "🏆 Monitor student learning through questions",
    "💡 Foster independent learning habits"
])

# Slide 15: Technical Implementation
add_content_slide(prs, "How It Works", [
    "1️⃣ User types a message in the chat interface",
    "2️⃣ Frontend sends message via API to backend",
    "3️⃣ Backend receives request and context",
    "4️⃣ Gemini AI processes with system prompt",
    "5️⃣ Response formatted professionally",
    "6️⃣ Frontend displays response in chat"
])

# Slide 16: System Requirements
add_content_slide(prs, "Setup Requirements", [
    "📋 Requirements:",
    "  • Node.js v18+ and npm",
    "  • Google Gemini API key (free tier available)",
    "  • Modern web browser (Chrome, Firefox, Safari, Edge)",
    "  • Internet connection",
    "---",
    "⏱️ Setup Time: ~5 minutes"
])

# Slide 17: Installation Guide
add_content_slide(prs, "Quick Start Guide", [
    "1. Clone the project or download files",
    "2. Run: npm install (install dependencies)",
    "3. Create .env file with GOOGLE_API_KEY",
    "4. Run: npm start (start server)",
    "5. Open: http://localhost:3003 in browser",
    "6. Start chatting! 🎉"
])

# Slide 18: How to Use
add_content_slide(prs, "Using the Chatbot", [
    "✏️ Type your question in the message box",
    "📤 Press 'Send' button or hit Enter",
    "⏳ Wait for AI response (usually <2 seconds)",
    "💬 Continue the conversation naturally",
    "🔄 Full chat history is maintained",
    "✨ Responses get better with context"
])

# Slide 19: Future Enhancements
add_content_slide(prs, "Future Roadmap", [
    "🔐 User authentication & individual chat histories",
    "📊 Analytics dashboard for educators",
    "🌐 Multi-language support",
    "🎯 Fine-tuned models for specific subjects",
    "📱 Mobile app (iOS/Android)",
    "🚀 Advanced features: document upload, PDF analysis"
])

# Slide 20: Conclusion
add_title_slide(prs, "Thank You!", "Questions? Contact or Try It Today!\nhttp://localhost:3003\n\nPowered by Google Gemini AI & Node.js")

# Save presentation
output_path = "c:/Users/babua/OneDrive/Desktop/mychaat/College_AI_Assistant.pptx"
prs.save(output_path)
print(f"✅ Presentation created successfully!")
print(f"📄 File saved to: {output_path}")
print(f"📊 Total slides: {len(prs.slides)}")
